from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def export_to_excel(videos: list[dict], output_path: str) -> None:
    df = pd.DataFrame(videos)

    # Format upload date
    if "upload_date" in df.columns:
        df["upload_date"] = pd.to_datetime(
            df["upload_date"],
            format="%Y%m%d",
            errors="coerce",
        )

        df = df.sort_values(
            by="upload_date",
            ascending=False,
        )

        # Keep date only
        df["upload_date"] = df["upload_date"].dt.date

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    # Export dataframe
    df.to_excel(
        output,
        index=False,
        sheet_name="Videos",
    )

    # Format workbook
    workbook = load_workbook(output)
    worksheet = workbook["Videos"]

    # Freeze header
    worksheet.freeze_panes = "A2"

    # Add filter
    worksheet.auto_filter.ref = worksheet.dimensions

    # Header style
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Column widths
    column_widths = {
        "A": 14,   # upload_date
        "B": 65,   # title
        "C": 100,  # description
        "D": 45,   # video_url
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # Wrap text and align cells
    for row in worksheet.iter_rows(min_row=2):
        row[0].alignment = Alignment(
            vertical="top",
            horizontal="center",
        )

        row[1].alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        row[2].alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        row[3].alignment = Alignment(
            vertical="top",
            wrap_text=False,
        )

        # Make URL clickable
        if row[3].value:
            row[3].hyperlink = row[3].value
            row[3].style = "Hyperlink"

    # Give each row enough height
    for row_number in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_number].height = 45

    workbook.save(output)

    print(f"\nExcel saved to: {output.resolve()}")